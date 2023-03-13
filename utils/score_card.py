import openpyxl
from datetime import date
import os
from openpyxl.formula.translate import Translator


from utils.helpers import use_dotenv

use_dotenv()


def extend_formulae(sheet, start_row: int, col_letter: str):
    last_row = sheet.max_row
    i = start_row
    while i < last_row:
        i += 1
        origin_formula = f"{col_letter}15"
        formula = sheet[origin_formula].value
        print(formula)
        sheet[f"{col_letter}{i}"] = Translator(
            formula, origin=origin_formula
        ).translate_formula(f"{col_letter}{i}")


def handle_scorecard(scorecard):

    print(f"SCORCARD DATA: {scorecard}")
    scorecard_total_pop = (
        scorecard["rep_lt_pop"]
        + scorecard["plant_nd_pop"]
        + scorecard["spk_plant_pop"]
        + scorecard["std_cost_pop"]
    )
    scorecard_total_isu = (
        scorecard["rep_lt_isu"]
        + scorecard["plant_nd_isu"]
        + scorecard["spk_plant_isu"]
        + scorecard["std_cost_isu"]
    )
    scorecard_percent = (scorecard_total_isu / scorecard_total_pop) * 100
    scorecard_percent = round(scorecard_percent, 2)
    print(
        f"Scorecard issues {scorecard_total_isu}\nScorecard populations: {scorecard_total_pop}\nScorecard percent: {scorecard_percent}"
    )

    f_scorecard = os.environ["SCORECARD"]
    wb_scorecard = openpyxl.load_workbook(f_scorecard)
    ws_ops_plan = wb_scorecard["Operations and Planning"]

    new_row = len(list(ws_ops_plan.rows)) + 1

    today = date.today().strftime("%d/%m/%y")
    date_my = date.today().strftime("%b-%y")
    fy = "FY" + date_my[-2:]

    ws_ops_plan[f"A{new_row}"] = date_my
    ws_ops_plan[f"B{new_row}"] = fy
    ws_ops_plan[f"P{new_row}"] = today

    ws_ops_plan[f"D{new_row}"] = scorecard["rep_lt_isu"]
    ws_ops_plan[f"E{new_row}"] = scorecard["rep_lt_pop"]
    ws_ops_plan[f"G{new_row}"] = scorecard["plant_nd_isu"]
    ws_ops_plan[f"H{new_row}"] = scorecard["plant_nd_pop"]
    ws_ops_plan[f"J{new_row}"] = scorecard["spk_plant_isu"]
    ws_ops_plan[f"K{new_row}"] = scorecard["spk_plant_pop"]
    ws_ops_plan[f"M{new_row}"] = scorecard["std_cost_isu"]
    ws_ops_plan[f"N{new_row}"] = scorecard["std_cost_pop"]

    tbl_ops_plan = ws_ops_plan.tables["Operations_and_Planning"]
    tbl_ops_plan.ref = f"A1:P{new_row}"

    extend_formulae(ws_ops_plan, 16, "C")

    wb_scorecard.save(f_scorecard)
